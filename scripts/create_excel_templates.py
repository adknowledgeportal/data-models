#!/usr/bin/env python3
# ------------------------------------------------------------------------------
# Title:        Build ADKP Excel templates
# Stage:        Data dictionary templates (not stage-specific)
# Author:       @grayjulia
# Description:  Python version of createExcelTemplates.R (adkp-data-ingress),
#               run automatically by .github/workflows/build-excel-templates.yml
#               after each release's schemas are registered
#
# Inputs:       registered sage.schemas.ad schemas; AD.model.csv (at the
#               release's tag)
# Outputs:      output/<Template>.xlsx, uploaded to the data dictionary folder
#
# Usage:        python scripts/create_excel_templates.py --folder-id synXXX [--version X] [--dry-run]
# Dependencies: scripts/requirements-excel-templates.in
# ------------------------------------------------------------------------------
"""Build the ADKP Excel templates from the registered sage.schemas.ad JSON
schemas and upload them to the data dictionary folder on Synapse.

Each template gets a Manifest sheet (one column per attribute, with dropdowns,
header notes, and required columns highlighted), a hidden ValidValues sheet
backing the dropdowns, and a DataDictionary sheet of every attribute/value
with its definition from AD.model.csv.

Only schemas registered at the release version are built. Schemas dropped
from the model stay registered at their last version and are skipped.

Differences from createExcelTemplates.R (the workbooks are otherwise the same,
cell for cell and in formatting):
  - builds only schemas at the release version; the R script builds every
    schema in the org, including ones removed from the model. Only the newest
    release can be uploaded; older ones can be built with --dry-run
  - fetches each schema at that exact version rather than the unversioned
    (latest) URL, and with --version waits for the release to be registered
  - lists and fetches schemas without a Synapse login; login is only needed
    to upload, which uses synapseclient instead of synapser and labels each
    file version with the model version (version comment, plus modelVersion
    and releaseDate annotations; the date is when the release was
    registered in sage.schemas.ad, in UTC)
  - exits with an error if any template fails to build or upload; the R
    script logs the failure and carries on
  - reads AD.model.csv at the release's tag rather than main, so
    definitions and column order match that release even if main has
    changed since
  - doesn't highlight an empty species cell on experimentalData rows: species
    is already required (and highlighted) on every row of the templates that
    have it
  - writes the dropdowns' warning error style directly, so there's no
    post-save XML patch
  - the DataDictionary filter covers the whole table, not just the header row
  - keeps each attribute's valid values as a list; the R script joins them
    with "; " and splits them again, which would break a value containing "; "
    (none do today)

Usage:
    python scripts/create_excel_templates.py --dry-run        # build only, newest registered version
    python scripts/create_excel_templates.py --version 1.0.9 --folder-id syn00000000

The upload folder comes from --folder-id or the TEMPLATE_FOLDER_ID environment
variable (a repository variable in the workflow).

Requires: pip install -r scripts/requirements-excel-templates.in
(the workflows install the exact versions in requirements-excel-templates.txt)

Auth (upload only): SYNAPSE_AUTH_TOKEN, or authtoken in ~/.synapseConfig.
"""

import argparse
import csv
import io
import os
import re
import sys
import time
from concurrent.futures import ThreadPoolExecutor
from dataclasses import dataclass
from pathlib import Path

import requests
import synapseclient
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from packaging.version import Version
from requests.adapters import HTTPAdapter
from synapseclient.models import File, Folder
from urllib3.util.retry import Retry

# =========================================================
# Settings
# =========================================================

REPO_URL = "https://repo-prod.prod.sagebase.org/repo/v1"
ORG_NAME = "sage.schemas.ad"
# the model as it was at the release's tag, so the definitions and column
# order match the schemas being built even after main moves on
MODEL_CSV_URL = "https://raw.githubusercontent.com/adknowledgeportal/data-models/refs/tags/v{version}/AD.model.csv"

MANIFEST_ROWS = 1000  # rows given borders, fills, and dropdowns
MAX_WORKERS = 5  # parallel requests to Synapse

# Schemas are read with direct REST calls rather than synapseclient's
# JSONSchema classes: those wrap the same endpoints and are still changing
# between releases, while the REST endpoints are Synapse's stable public API.
# This session retries dropped connections and busy/error responses up to 5
# times (at once, then after 2, 4, 8 and 16 seconds). The POSTs are read-only
# list queries, so they're safe to retry too.
SESSION = requests.Session()
SESSION.mount("https://", HTTPAdapter(
    max_retries=Retry(total=5, backoff_factor=1, status_forcelist=(429, 500, 502, 503, 504),
                      allowed_methods=None, respect_retry_after_header=True),
    pool_maxsize=MAX_WORKERS))

# Workbook styles. Fonts are set on every styled cell: a cell given any style
# starts from openpyxl's Calibri 11, not the workbook's default font.
ARIAL = Font(name="Arial", size=10)
ARIAL_BOLD = Font(name="Arial", size=10, bold=True)
ARIAL_BOLD_WHITE = Font(name="Arial", size=10, bold=True, color="FFFFFF")
MEDIUM_GREY_SIDE = Side(style="medium", color="E2E3E3")
THIN_SIDE = Side(style="thin")
MANIFEST_BORDER = Border(left=MEDIUM_GREY_SIDE, right=MEDIUM_GREY_SIDE,
                         top=MEDIUM_GREY_SIDE, bottom=MEDIUM_GREY_SIDE)
DD_BORDER = Border(left=THIN_SIDE, right=THIN_SIDE, top=THIN_SIDE, bottom=THIN_SIDE)
REQUIRED_FILL = PatternFill(fill_type="solid", start_color="B5E3E8", end_color="B5E3E8")
OPTIONAL_HEADER_FILL = PatternFill(fill_type="solid", start_color="E0E0E0", end_color="E0E0E0")
DD_HEADER_FILL = PatternFill(fill_type="solid", start_color="375623", end_color="375623")
DD_BAND_FILL = PatternFill(fill_type="solid", start_color="EBF3E8", end_color="EBF3E8")
DD_COLUMN_WIDTHS = {"A": 29.41, "B": 41.41, "C": 33.41, "D": 61.41}

VALIDATION_ERROR_TITLE = "Value not in standard list"
VALIDATION_ERROR = "This entry isn't one of the standard values for this field. Keep it anyway?"
MULTISELECT_NOTE = ("\nMultiple values allowed: type your selections separated by \",\" with no space "
                    "(e.g. \"Human,Mouse\"). Excel will show a warning since it isn't a single "
                    "listed value, but the entry is accepted.")


# =========================================================
# 1. Registered schemas
# =========================================================

@dataclass(frozen=True)
class Attribute:
    """One template column, from a property of the registered JSON schema."""
    name: str
    description: str | None
    type: str | None  # JSON schema type, e.g. "string" or "array"
    valid_values: list[str] | None  # None when any value is allowed
    required: bool


def post_paginated(path: str, body: dict) -> list[dict]:
    """All items from a paginated Synapse POST endpoint (public, no login)."""
    items = []
    while True:
        resp = SESSION.post(REPO_URL + path, json=body, timeout=60)
        resp.raise_for_status()
        data = resp.json()
        items += data.get("page", [])
        if not data.get("nextPageToken"):
            return items
        body = {**body, "nextPageToken": data["nextPageToken"]}


def registered_versions() -> dict[str, dict[str, str]]:
    """{schema name: {semantic version: registration time (ISO 8601, UTC)}}, oldest first"""
    names = [s["schemaName"] for s in post_paginated("/schema/list", {"organizationName": ORG_NAME})]

    def versions(name: str) -> dict[str, str]:
        found = post_paginated("/schema/version/list",
                               {"organizationName": ORG_NAME, "schemaName": name})
        return {v["semanticVersion"]: v["createdOn"]
                for v in sorted(found, key=lambda v: v["createdOn"]) if v.get("semanticVersion")}

    with ThreadPoolExecutor(max_workers=MAX_WORKERS) as pool:
        return dict(zip(names, pool.map(versions, names)))


def resolve_release(version: str | None,
                    wait_minutes: float) -> tuple[str, list[str], str, str]:
    """Pick the release version and the schemas registered at it.

    With no version, the newest registered one is used. With a version, polls
    until at least one schema has it, since this may run right after the
    data-models release workflow registers them. Returns the release, the
    schemas registered at it, the newest version registered overall, and the
    date (UTC) the release was registered in the org.
    """
    deadline = time.monotonic() + wait_minutes * 60
    while True:
        registered = registered_versions()
        newest = max((v for vs in registered.values() for v in vs), key=Version)
        release = version or newest
        in_release = sorted(n for n, vs in registered.items() if release in vs)
        if in_release:
            break
        if time.monotonic() >= deadline:
            sys.exit(f"No {ORG_NAME} schemas registered at {release} after {wait_minutes} minute(s)")
        print(f"No {ORG_NAME} schemas at {release} yet, checking again in 30s", flush=True)
        time.sleep(30)

    for name, vs in sorted(registered.items()):
        if release not in vs:
            print(f"[{name}] skipped: not registered at {release} "
                  f"(newest: {list(vs)[-1] if vs else 'none'})")
    # when the release's schemas were registered in the production org; this
    # is the production release date even if the GitHub release started as a
    # pre-release. Times are UTC.
    registered_on = min(registered[n][release] for n in in_release)[:10]
    print(f"Building {len(in_release)} template(s) from {ORG_NAME} {release} "
          f"(registered {registered_on} UTC)")
    return release, in_release, newest, registered_on


def fetch_schema(name: str, version: str) -> dict:
    resp = SESSION.get(f"{REPO_URL}/schema/type/registered/{ORG_NAME}-{name}-{version}", timeout=60)
    resp.raise_for_status()
    return resp.json()


def as_text(value: object) -> str:
    # match how R wrote enum values: TRUE/FALSE, and 1 rather than 1.0
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def extract_attrs(schema: dict) -> list[Attribute]:
    """One Attribute per property, in schema order."""
    required = set(schema.get("required", []))
    attrs = []
    for name, prop in (schema.get("properties") or {}).items():
        # array/multiselect properties carry their allowed values under items.enum,
        # not enum directly on the property -- that's where enum lives for singleselect
        enum = prop.get("enum") or (prop.get("items") or {}).get("enum")
        attrs.append(Attribute(
            name=name,
            description=prop.get("description"),
            type=prop.get("type"),
            valid_values=[as_text(v) for v in enum] if enum else None,
            required=name in required,
        ))
    return attrs


# =========================================================
# 2. Definitions and column order from AD.model.csv
# =========================================================

Definitions = dict[tuple[str | None, str | None], str | None]  # (Property, Value) -> Definition
ColumnOrder = dict[str, list[str]]  # normalized template name -> DependsOn attributes


def norm_name(x: str) -> str:
    return re.sub(r"[^a-z0-9]", "", x.lower())


def load_model(source: str) -> tuple[Definitions, ColumnOrder]:
    """(definitions {(Property, Value): Definition}, order {template key: [attributes]})

    source is a URL or a local path to AD.model.csv.
    """
    if re.match(r"https?://", source):
        resp = SESSION.get(source, timeout=60)
        resp.raise_for_status()
        content = resp.content
    else:
        content = Path(source).read_bytes()
    rows = list(csv.DictReader(io.StringIO(content.decode("utf-8-sig"))))

    def cell(row: dict[str, str], col: str) -> str | None:
        # blank and "NA" are missing, as they were for readr
        value = (row.get(col) or "").strip()
        return None if value in ("", "NA") else value

    definitions, order = {}, {}
    for row in rows:
        definitions.setdefault((cell(row, "Parent"), cell(row, "Attribute")),
                               cell(row, "Description"))
        depends_on = cell(row, "DependsOn")
        if (cell(row, "IsTemplate") or "").upper() == "TRUE" and depends_on:
            order.setdefault(norm_name(row["Attribute"]),
                             [a.strip() for a in re.split(r",\s*", depends_on)])
    return definitions, order


def reorder_by_dependson(attrs: list[Attribute], template_name: str,
                         order: ColumnOrder) -> list[Attribute]:
    """Reorder attributes to match the template's DependsOn order."""
    ordered = order.get(norm_name(template_name))
    if ordered is None:
        print(f"[{template_name}] no DependsOn order found in data model — keeping schema order")
        return attrs

    by_name = {a.name: a for a in attrs}
    matched = [a for a in ordered if a in by_name]  # in DependsOn order
    leftover = [a.name for a in attrs if a.name not in matched]
    if leftover:
        print(f"[{template_name}] {len(leftover)} attribute(s) not in DependsOn, "
              f"appended at end: {', '.join(leftover)}")
    return [by_name[a] for a in dict.fromkeys(matched + leftover)]


# =========================================================
# 3. Workbook
# =========================================================

def set_default_font(wb: Workbook) -> None:
    """Make Arial 10 the workbook's default font (its "Normal" style).

    openpyxl has no public way to do this, so this uses its private
    _named_styles attribute. If an openpyxl update removes or renames it,
    test_default_font_is_arial in tests/test_create_excel_templates.py fails;
    see docs/excel-templates-maintenance.md for what to do.
    """
    wb._named_styles["Normal"].font = Font(name="Arial", size=10)


def build_manifest(wb: Workbook, attrs: list[Attribute]) -> None:
    ws = wb.active
    ws.title = "Manifest"
    values_ws = wb.create_sheet("ValidValues")
    values_ws.sheet_state = "hidden"

    for col, attr in enumerate(attrs, start=1):
        letter = get_column_letter(col)
        header = ws.cell(row=1, column=col, value=attr.name)
        header.font = ARIAL_BOLD
        header.fill = REQUIRED_FILL if attr.required else OPTIONAL_HEADER_FILL
        for row in range(1, MANIFEST_ROWS + 1):
            cell = ws.cell(row=row, column=col)
            cell.border = MANIFEST_BORDER
            if row > 1:
                cell.font = ARIAL
            if attr.required:
                cell.fill = REQUIRED_FILL

        # dropdown backed by this attribute's column in the hidden sheet; a
        # warning (not a stop) so multiselect entries and new values are allowed
        values = attr.valid_values
        if values:
            for row, value in enumerate(values, start=1):
                values_ws.cell(row=row, column=col, value=value).font = ARIAL
            dv = DataValidation(type="list", formula1=f"ValidValues!${letter}$1:${letter}${len(values)}",
                                allow_blank=True, showInputMessage=True, showErrorMessage=True,
                                errorStyle="warning", errorTitle=VALIDATION_ERROR_TITLE,
                                error=VALIDATION_ERROR)
            dv.add(f"{letter}2:{letter}{MANIFEST_ROWS}")
            ws.add_data_validation(dv)

        if attr.description:
            text = attr.description
            if attr.type == "array" and values:
                text += MULTISELECT_NOTE
            # size the note box to fit the text -- the default truncates
            # anything longer than a short one-liner
            lines = max(2, -(-len(text) // 45) + 1)
            header.comment = Comment(text, "", width=3 * 64, height=lines * 20)


def build_data_dictionary(wb: Workbook, attrs: list[Attribute], definitions: Definitions,
                          template_name: str) -> None:
    ws = wb.create_sheet("DataDictionary")
    rows = [(a.name, a.description, value, definitions.get((a.name, value)))
            for a in attrs for value in (a.valid_values or [None])]

    unmatched = {(p, v) for p, _, v, d in rows if v is not None and d is None}
    if unmatched:
        print(f"[{template_name}] {len(unmatched)} value(s) with no definition match")

    ws.append(["Property", "Description", "Value", "Definition"])
    for row in rows:
        ws.append(row)

    header_alignment = Alignment(horizontal="left")
    body_alignment = Alignment(wrap_text=True, vertical="top")

    # banded shading per attribute group
    group, previous = 0, object()
    for r, (prop, *_) in enumerate(rows, start=2):
        if prop != previous:
            group, previous = group + 1, prop
        for c in range(1, 5):
            cell = ws.cell(row=r, column=c)
            cell.font = ARIAL
            cell.border = DD_BORDER
            cell.alignment = body_alignment
            if group % 2 == 1:
                cell.fill = DD_BAND_FILL
    for c in range(1, 5):
        cell = ws.cell(row=1, column=c)
        cell.border = DD_BORDER
        cell.font = ARIAL_BOLD_WHITE
        cell.fill = DD_HEADER_FILL
        cell.alignment = header_alignment

    for letter, width in DD_COLUMN_WIDTHS.items():
        ws.column_dimensions[letter].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:D{len(rows) + 1}"


def build_template_workbook(attrs: list[Attribute], template_name: str, definitions: Definitions,
                            order: ColumnOrder, out_dir: str | Path) -> Path:
    attrs = reorder_by_dependson(attrs, template_name, order)

    wb = Workbook()
    set_default_font(wb)
    build_manifest(wb, attrs)
    build_data_dictionary(wb, attrs, definitions, template_name)

    out_path = Path(out_dir) / f"{template_name}.xlsx"
    wb.save(out_path)
    return out_path


# =========================================================
# 4. Upload to Synapse — new version of an existing file with the
#    same name (case-insensitive), or a new file
# =========================================================

def upload(paths: list[Path], release: str, registered_on: str, folder_id: str | None,
           dry_run: bool) -> list[str]:
    if dry_run and not folder_id:
        for path in paths:
            print(f"[{path.stem}] built (no --folder-id, so not checked against Synapse)")
        return []

    syn = synapseclient.Synapse()
    syn.login(silent=True)

    folder = Folder(id=folder_id).sync_from_synapse(download_file=False, recursive=False)
    existing = {f.name.lower(): f.id for f in folder.files}

    # every file version records the model version it was built from; the
    # file names stay unversioned so each template keeps its Synapse ID
    labels = {"modelVersion": [release], "releaseDate": [registered_on]}
    print(f"Labeling uploads: {ORG_NAME} {release}, " +
          ", ".join(f"{k}={v[0]}" for k, v in labels.items()))

    failed = []
    for path in paths:
        existing_id = existing.get(path.name.lower())
        if dry_run:
            print(f"[{path.stem}] would {'update ' + existing_id if existing_id else 'create'}")
            continue
        try:
            if existing_id:
                entity = File(id=existing_id, download_file=False).get()
                entity.path = str(path)
            else:
                entity = File(path=str(path), name=path.name, parent_id=folder_id)
            entity.version_comment = f"{ORG_NAME} {release}"
            entity.annotations = {**(entity.annotations or {}), **labels}
            entity = entity.store()
            print(f"[{path.stem}] stored as {entity.id} (v{entity.version_number})")
        except Exception as e:
            print(f"[{path.stem}] upload failed: {e}")
            failed.append(path.stem)
    return failed


def _try(fn, *args):
    """fn(*args), or the exception it raised, so one failure in a parallel
    map doesn't stop the rest."""
    try:
        return fn(*args)
    except Exception as e:
        return e


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    parser.add_argument("--version", help="release to build, e.g. 1.0.9 (default: newest registered)")
    parser.add_argument("--wait-minutes", type=float, default=10,
                        help="how long to wait for --version to be registered (default: 10)")
    parser.add_argument("--out-dir", default="output")
    parser.add_argument("--folder-id", default=os.environ.get("TEMPLATE_FOLDER_ID"),
                        help="Synapse folder to upload to (default: $TEMPLATE_FOLDER_ID)")
    parser.add_argument("--model-csv",
                        help="AD.model.csv URL or local path (default: the release's tag "
                             "on GitHub)")
    parser.add_argument("--dry-run", action="store_true", help="build, but don't upload")
    args = parser.parse_args()
    if not args.folder_id and not args.dry_run:
        parser.error("--folder-id (or TEMPLATE_FOLDER_ID) is required unless --dry-run")

    release, names, newest, registered_on = resolve_release(
        args.version and args.version.lstrip("v"), args.wait_minutes)
    # the folder holds one file per template, so uploading an older release
    # would replace the current templates with older ones
    if release != newest and not args.dry_run:
        sys.exit(f"{release} isn't the newest registered version ({newest}); uploading it "
                 "would replace the current templates. Use --dry-run to build it locally.")
    model_csv = args.model_csv or MODEL_CSV_URL.format(version=release)
    print(f"Reading definitions and column order from {model_csv}")
    try:
        definitions, order = load_model(model_csv)
    except Exception as e:
        sys.exit(f"Couldn't read AD.model.csv from {model_csv} ({e}). "
                 "Pass --model-csv to use a different copy.")
    Path(args.out_dir).mkdir(exist_ok=True)

    with ThreadPoolExecutor(max_workers=MAX_WORKERS) as pool:
        schemas = dict(zip(names, pool.map(lambda n: _try(fetch_schema, n, release), names)))

    # one bad schema shouldn't stop the rest, but the run still fails at the end
    built, not_built = [], []
    for name in names:
        try:
            if isinstance(schemas[name], Exception):
                raise schemas[name]
            attrs = extract_attrs(schemas[name])
            if not attrs:
                raise ValueError("schema has no properties")
            built.append(build_template_workbook(attrs, name, definitions, order, args.out_dir))
        except Exception as e:
            print(f"[{name}] failed: {e}")
            not_built.append(name)

    not_stored = upload(built, release, registered_on, args.folder_id, args.dry_run)
    if not_built or not_stored:
        sys.exit(f"{len(not_built)} template(s) not built ({', '.join(not_built)}), "
                 f"{len(not_stored)} not stored ({', '.join(not_stored)})")


if __name__ == "__main__":
    main()
