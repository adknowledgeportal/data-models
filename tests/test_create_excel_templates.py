"""Tests for scripts/create_excel_templates.py.

These run on every pull request that changes the script, its packages or
these tests (.github/workflows/test-excel-templates.yml), including
Dependabot's monthly package updates. If one fails, see
docs/excel-templates-maintenance.md.

Run them locally from the repository folder with:  python -m pytest tests -v
"""
import dataclasses
import xml.etree.ElementTree as ET
import zipfile

import pytest
from openpyxl import load_workbook
from synapseclient.models import File, Folder

import create_excel_templates as cet

MAINTENANCE_DOC = "docs/excel-templates-maintenance.md"

# registration times are ISO 8601 in UTC, oldest first, as Synapse returns them
REGISTERED = {
    "IndividualTemplate": {"1.0.8": "2026-01-05T20:00:00.000Z", "1.0.9": "2026-02-02T23:30:00.000Z"},
    "BiospecimenTemplate": {"1.0.9": "2026-02-03T00:10:00.000Z"},
    "RetiredTemplate": {"1.0.8": "2026-01-05T20:01:00.000Z"},
}


def attribute(name, required=False, valid_values=None, type="string", description="A description"):
    return cet.Attribute(name=name, description=description, type=type,
                         valid_values=valid_values, required=required)


# ---------------------------------------------------------
# Choosing the release and its schemas
# ---------------------------------------------------------

def test_newest_release_builds_only_schemas_registered_at_it(monkeypatch):
    monkeypatch.setattr(cet, "registered_versions", lambda: REGISTERED)
    release, names, newest, registered_on = cet.resolve_release(None, wait_minutes=0)
    assert release == newest == "1.0.9"
    assert names == ["BiospecimenTemplate", "IndividualTemplate"]  # not RetiredTemplate
    # the earliest registration of the release, as a UTC date
    assert registered_on == "2026-02-02"


def test_older_release_builds_every_schema_registered_at_it(monkeypatch):
    # an older release includes schemas that have newer versions since
    monkeypatch.setattr(cet, "registered_versions", lambda: REGISTERED)
    release, names, newest, _ = cet.resolve_release("1.0.8", wait_minutes=0)
    assert release == "1.0.8"
    assert newest == "1.0.9"
    assert names == ["IndividualTemplate", "RetiredTemplate"]


def test_versions_compare_as_numbers_not_text(monkeypatch):
    # as text, "1.0.9" > "1.0.10"
    monkeypatch.setattr(cet, "registered_versions", lambda: {
        "IndividualTemplate": {"1.0.9": "2026-01-01T00:00:00.000Z", "1.0.10": "2026-02-01T00:00:00.000Z"},
    })
    _, _, newest, _ = cet.resolve_release(None, wait_minutes=0)
    assert newest == "1.0.10"


def test_uploading_an_older_release_is_refused(monkeypatch):
    # the folder holds one file per template, so this would replace the
    # current templates with older ones
    monkeypatch.setattr(cet, "registered_versions", lambda: REGISTERED)
    monkeypatch.setattr("sys.argv", ["create_excel_templates.py", "--version", "1.0.8",
                                     "--folder-id", "syn00000000"])
    with pytest.raises(SystemExit, match="isn't the newest registered version"):
        cet.main()


def test_resolve_only_reports_the_release_without_building(monkeypatch, tmp_path):
    # the workflow uses this to pick one version for the build and the
    # release-page status, so they can't disagree
    output = tmp_path / "github_output"
    monkeypatch.setenv("GITHUB_OUTPUT", str(output))
    monkeypatch.delenv("TEMPLATE_FOLDER_ID", raising=False)
    monkeypatch.setattr(cet, "registered_versions", lambda: REGISTERED)
    monkeypatch.setattr(cet, "load_model", lambda source: pytest.fail("built templates"))
    monkeypatch.setattr("sys.argv", ["create_excel_templates.py", "--resolve-only"])
    cet.main()
    assert output.read_text() == "release=1.0.9\nregistered_on=2026-02-02\n"


def test_error_annotations_escape_external_text(monkeypatch, capsys):
    # error text from Synapse or HTTP can contain newlines or %, which must not
    # end the GitHub workflow command early or start a new one
    monkeypatch.setenv("GITHUB_ACTIONS", "true")
    with pytest.raises(SystemExit):
        cet.fail("Couldn't read the model (100% failed\n::warning::injected)")
    cet.report_error("IndividualTemplate", "upload failed: line one\r\nline two")
    lines = capsys.readouterr().out.splitlines()
    assert lines[0] == ("::error title=Excel templates::Couldn't read the model "
                        "(100%25 failed%0A::warning::injected)")
    assert lines[-1] == "::error title=IndividualTemplate::upload failed: line one%0D%0Aline two"
    assert not any(line.startswith("::warning::") for line in lines)


# ---------------------------------------------------------
# Reading the schema and AD.model.csv
# ---------------------------------------------------------

def test_extract_attrs_reads_properties_in_schema_order():
    schema = {
        "required": ["species"],
        "properties": {
            "species": {"description": "Species", "type": "string", "enum": ["Human", "Mouse"]},
            "assay": {"description": "Assay", "type": "array", "items": {"enum": ["RNAseq"]}},
            "isPostMortem": {"type": "boolean", "enum": [True, False]},
            "count": {"type": "number", "enum": [1.0, 2.5]},
            "notes": {"type": "string"},
        },
    }
    attrs = cet.extract_attrs(schema)
    assert [a.name for a in attrs] == ["species", "assay", "isPostMortem", "count", "notes"]
    assert attrs[0].required and not attrs[1].required
    assert attrs[1].valid_values == ["RNAseq"]  # multiselect values are under items.enum
    assert attrs[2].valid_values == ["TRUE", "FALSE"]  # as the R script wrote them
    assert attrs[3].valid_values == ["1", "2.5"]
    assert attrs[4].valid_values is None


def test_load_model_reads_definitions_and_column_order(tmp_path):
    csv_path = tmp_path / "AD.model.csv"
    csv_path.write_text(
        "Attribute,Description,Parent,DependsOn,IsTemplate\n"
        "IndividualTemplate,Template,NA,\"individualID, species\",TRUE\n"
        "species,Species of the individual,NA,,FALSE\n"
        "Human,A human,species,,FALSE\n"
        "Mouse,NA,species,,FALSE\n",
        encoding="utf-8")
    definitions, order = cet.load_model(str(csv_path))
    assert definitions[("species", "Human")] == "A human"
    assert definitions[("species", "Mouse")] is None  # "NA" means missing
    assert order == {"individualtemplate": ["individualID", "species"]}


def test_columns_follow_dependson_order_with_extras_at_the_end():
    attrs = [attribute("species"), attribute("notes"), attribute("individualID")]
    order = {"individualtemplate": ["individualID", "species", "notInSchema"]}
    reordered = cet.reorder_by_dependson(attrs, "IndividualTemplate", order)
    assert [a.name for a in reordered] == ["individualID", "species", "notes"]


def test_columns_keep_schema_order_when_template_not_in_model():
    attrs = [attribute("b"), attribute("a")]
    assert cet.reorder_by_dependson(attrs, "UnknownTemplate", {}) == attrs


# ---------------------------------------------------------
# The workbook
# ---------------------------------------------------------

@pytest.fixture
def workbook_path(tmp_path):
    attrs = [
        attribute("individualID", required=True),
        attribute("assay", type="array", valid_values=["RNAseq", "ATACseq"]),
        attribute("notes", description=None),
    ]
    definitions = {("assay", "RNAseq"): "RNA sequencing"}
    return cet.build_template_workbook(attrs, "IndividualTemplate", definitions, {}, tmp_path)


def test_default_font_is_arial(workbook_path):
    """Fails if an openpyxl update breaks set_default_font, which uses a
    private openpyxl attribute. See the maintenance doc."""
    with zipfile.ZipFile(workbook_path) as z:
        styles = ET.fromstring(z.read("xl/styles.xml"))
    # the default font is the one used by the "Normal" cell style:
    # cellStyle Normal -> its cellStyleXfs entry -> that entry's font
    ns = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    normal = styles.find("m:cellStyles/m:cellStyle[@name='Normal']", ns)
    style_xf = styles.findall("m:cellStyleXfs/m:xf", ns)[int(normal.get("xfId"))]
    font = styles.findall("m:fonts/m:font", ns)[int(style_xf.get("fontId"))]
    name, size = font.find("m:name", ns).get("val"), font.find("m:sz", ns).get("val")
    assert (name, size) == ("Arial", "10"), (
        f"The workbook's default font is {name} {size}, not Arial 10, so set_default_font() "
        f"stopped working, probably after an openpyxl update. See {MAINTENANCE_DOC}.")


def test_set_default_font_still_works_with_this_openpyxl():
    """set_default_font uses openpyxl's private _named_styles attribute.
    Fails with a clear message if an openpyxl update removes it."""
    try:
        cet.set_default_font(cet.Workbook())
    except (AttributeError, KeyError, TypeError) as e:
        pytest.fail(f"set_default_font() no longer works with this openpyxl version ({e!r}). "
                    f"See {MAINTENANCE_DOC}.")


def test_manifest_sheet(workbook_path):
    wb = load_workbook(workbook_path)
    assert wb.sheetnames == ["Manifest", "ValidValues", "DataDictionary"]
    assert wb["ValidValues"].sheet_state == "hidden"

    ws = wb["Manifest"]
    assert [c.value for c in ws[1]] == ["individualID", "assay", "notes"]
    assert ws["A1"].fill.start_color.rgb.endswith("B5E3E8")  # required
    assert ws["B1"].fill.start_color.rgb.endswith("E0E0E0")  # optional
    assert ws["A500"].fill.start_color.rgb.endswith("B5E3E8")  # whole required column

    # dropdown from the hidden sheet, warning (not stop) so multiselect is allowed
    [dv] = ws.data_validations.dataValidation
    assert dv.formula1 == "ValidValues!$B$1:$B$2"
    assert dv.errorStyle == "warning"
    assert str(dv.sqref) == f"B2:B{cet.MANIFEST_ROWS}"

    assert cet.MULTISELECT_NOTE.strip() in ws["B1"].comment.text
    assert ws["C1"].comment is None  # no description, no note


def test_data_dictionary_sheet(workbook_path):
    ws = load_workbook(workbook_path)["DataDictionary"]
    rows = [[c.value for c in row] for row in ws.iter_rows()]
    assert rows[0] == ["Property", "Description", "Value", "Definition"]
    assert rows[1:] == [
        ["individualID", "A description", None, None],
        ["assay", "A description", "RNAseq", "RNA sequencing"],
        ["assay", "A description", "ATACseq", None],
        ["notes", None, None, None],
    ]
    assert ws.auto_filter.ref == "A1:D5"


# ---------------------------------------------------------
# synapseclient (upload)
# ---------------------------------------------------------

def test_synapseclient_still_has_what_upload_uses():
    """The upload step can't run in tests (it needs a Synapse login), so check
    that the synapseclient classes it uses still have the same names. Fails
    if a synapseclient update renames them. See the maintenance doc."""
    file_fields = {f.name for f in dataclasses.fields(File)}
    missing = {"id", "name", "path", "parent_id", "version_comment", "version_number",
               "annotations", "download_file"} - file_fields
    missing |= {f"File.{m}" for m in ("get", "store") if not callable(getattr(File, m, None))}
    missing |= {f"Folder.{m}" for m in ("sync_from_synapse",) if not callable(getattr(Folder, m, None))}
    if "files" not in {f.name for f in dataclasses.fields(Folder)}:
        missing.add("Folder.files")
    assert not missing, (f"synapseclient no longer has {sorted(missing)}, which upload() in "
                         f"create_excel_templates.py uses. See {MAINTENANCE_DOC}.")
